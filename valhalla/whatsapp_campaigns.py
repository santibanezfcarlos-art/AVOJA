"""Private historical contacts and assisted WhatsApp campaigns; never sends messages."""
import base64
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from urllib.parse import urlencode

DEFAULT_MESSAGE = ('Hola, {nombre},\n\n¡Tenemos una noticia! Ya está disponible la nueva plataforma de Valhalla Office: https://valhallaoffice.cl\n\nPuedes revisar disponibilidad, comprar horas y reservar tu box. Para reservar para el día siguiente, puedes hacerlo hasta las 21:30 del día anterior. El código de acceso se envía automáticamente por correo a las 22:00 del día anterior.\n\n¡Te invitamos a conocerla!\nEquipo Valhalla Office\n\nSi prefieres no recibir estas novedades, avísanos por aquí.')

def phone_number(value):
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) == 9 and digits.startswith('9'):
        digits = '56' + digits
    return digits if re.fullmatch(r'569\d{8}', digits) else ''

def read_contacts(encoded):
    from openpyxl import load_workbook
    if not isinstance(encoded, str) or len(encoded) > 4_000_000:
        raise ValueError('Selecciona un XLSX de hasta 3 MB.')
    try:
        raw = base64.b64decode(encoded, validate=True)
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 200 or sum(i.file_size for i in archive.infolist()) > 20_000_000:
                raise ValueError('Excel demasiado grande.')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.active
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v or '').strip() for v in next(rows)]
        required = ['Nombre', 'Apellido', 'Correo electrónico', 'Número de teléfono']
        if not all(h in headers for h in required):
            raise ValueError('Se requieren Nombre, Apellido, Correo electrónico y Número de teléfono.')
        records = {}; issues = []; total = 0; duplicates = 0
        for number, row in enumerate(rows, 2):
            if number > 5002: raise ValueError('Máximo 5.000 contactos por archivo.')
            if not any(v is not None for v in row): continue
            total += 1
            def cell(key):
                index = headers.index(key)
                return str(row[index] or '').strip() if index < len(row) else ''
            name = re.sub(r'^[\W_]+', '', cell('Nombre')).strip()
            name = ' '.join(name.split())
            full_name = (name + ' ' + cell('Apellido')).strip()[:200]
            phone = phone_number(cell('Número de teléfono'))
            email = cell('Correo electrónico').lower()[:200]
            if not phone or not name:
                issues.append({'row': number, 'name': full_name, 'reason': 'Revisar nombre o teléfono móvil chileno'})
                continue
            if phone in records: duplicates += 1
            record = records.setdefault(phone, {'phone': phone, 'name': full_name, 'first_name': name.split()[0].capitalize(), 'emails': []})
            if email and email not in record['emails']: record['emails'].append(email)
        return list(records.values()), {'rows': total, 'unique': len(records), 'duplicates': duplicates, 'issues': issues}
    except (zipfile.BadZipFile, OSError, KeyError, StopIteration) as error:
        raise ValueError('No se pudo leer el Excel.') from error
    finally:
        if 'workbook' in locals(): workbook.close()

def initialize(c):
    c.executescript('''
    CREATE TABLE IF NOT EXISTS whatsapp_contacts (
      phone TEXT PRIMARY KEY, name TEXT NOT NULL, first_name TEXT NOT NULL,
      emails TEXT NOT NULL, opted_out INTEGER NOT NULL DEFAULT 0, imported_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS whatsapp_campaigns (
      id INTEGER PRIMARY KEY, name TEXT NOT NULL, message TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS whatsapp_progress (
      campaign_id INTEGER NOT NULL REFERENCES whatsapp_campaigns(id), phone TEXT NOT NULL REFERENCES whatsapp_contacts(phone),
      status TEXT NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(campaign_id,phone));
    ''')

def report(store):
    with store.connection() as c:
        users = {r['email'].strip().lower() for r in c.execute("SELECT email FROM users WHERE role='professional' AND deleted_at=''")}
        contacts = []
        for row in c.execute('SELECT * FROM whatsapp_contacts ORDER BY name COLLATE NOCASE,phone'):
            record = dict(row); record['emails'] = json.loads(record['emails'])
            record['registered'] = any(e in users for e in record['emails'])
            contacts.append(record)
        return {'contacts': contacts, 'campaigns': [dict(r) for r in c.execute('SELECT * FROM whatsapp_campaigns ORDER BY id DESC')],
                'progress': [dict(r) for r in c.execute('SELECT * FROM whatsapp_progress')], 'default_message': DEFAULT_MESSAGE,
                'matching_note': 'Cruce por correo exacto; no se presume coincidencia por nombre.'}

def action(store, actor, body, operation):
    timestamp = datetime.now(timezone.utc).isoformat()
    if operation in ('preview', 'import'):
        records, info = read_contacts(body.get('content'))
        with store.connection(write=operation=='import') as c:
            known = {r['email'].lower() for r in c.execute("SELECT email FROM users WHERE role='professional' AND deleted_at=''")}
            info['registered'] = sum(any(e in known for e in r['emails']) for r in records)
            info['existing_contacts'] = sum(bool(c.execute('SELECT 1 FROM whatsapp_contacts WHERE phone=?',(r['phone'],)).fetchone()) for r in records)
            if operation == 'import':
                if body.get('confirmed') is not True: raise ValueError('Confirma la importación después de revisar la vista previa.')
                for r in records:
                    old = c.execute('SELECT emails FROM whatsapp_contacts WHERE phone=?',(r['phone'],)).fetchone()
                    emails = sorted(set(r['emails'] + (json.loads(old['emails']) if old else [])))
                    c.execute('''INSERT INTO whatsapp_contacts(phone,name,first_name,emails,imported_at) VALUES(?,?,?,?,?)
                      ON CONFLICT(phone) DO UPDATE SET name=excluded.name,first_name=excluded.first_name,emails=excluded.emails''',
                      (r['phone'],r['name'],r['first_name'],json.dumps(emails),timestamp))
                store.log(c,actor,'whatsapp_contacts_imported',len(records))
        return info | {'sample': records[:5]}
    with store.connection(write=True) as c:
        if operation == 'campaign':
            name = str(body.get('name') or '').strip(); message = str(body.get('message') or '').strip()
            if not 1 <= len(name) <= 100 or not 1 <= len(message) <= 3000: raise ValueError('Revisa el título y mensaje (máximo 3.000 caracteres).')
            cid = c.execute('INSERT INTO whatsapp_campaigns(name,message,created_at) VALUES(?,?,?)',(name,message,timestamp)).lastrowid
            store.log(c,actor,'whatsapp_campaign_created',cid)
            return {'id':cid}
        phone = body.get('phone')
        contact = c.execute('SELECT * FROM whatsapp_contacts WHERE phone=?',(phone,)).fetchone()
        if not contact: raise ValueError('Contacto no encontrado.')
        if operation == 'optout':
            c.execute('UPDATE whatsapp_contacts SET opted_out=1 WHERE phone=?',(phone,))
            store.log(c,actor,'whatsapp_contact_opted_out',phone)
            return {'ok':True}
        campaign = c.execute('SELECT * FROM whatsapp_campaigns WHERE id=?',(body.get('campaign_id'),)).fetchone()
        if not campaign: raise ValueError('Selecciona una campaña guardada.')
        if contact['opted_out']: raise ValueError('Este contacto está dado de baja.')
        status = body.get('status')
        if operation == 'open': status = 'opened'
        if status not in ('opened','sent_manual','skipped','pending'): raise ValueError('Estado inválido.')
        old = c.execute('SELECT status FROM whatsapp_progress WHERE campaign_id=? AND phone=?',(campaign['id'],phone)).fetchone()
        if operation == 'open' and old and old['status']=='sent_manual':
            raise ValueError('Ya está marcado como enviado. Restablece a pendiente si necesitas abrirlo otra vez.')
        c.execute('INSERT INTO whatsapp_progress VALUES(?,?,?,?) ON CONFLICT(campaign_id,phone) DO UPDATE SET status=excluded.status,updated_at=excluded.updated_at',
                  (campaign['id'],phone,status,timestamp))
        store.log(c,actor,'whatsapp_'+status,phone)
        message = campaign['message'].replace('{nombre}',contact['first_name'])
        return {'ok':True,'url':'https://wa.me/'+phone+'?'+urlencode({'text':message})} if operation=='open' else {'ok':True}
