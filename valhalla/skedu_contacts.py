"""Read Skedu's second worksheet for dated access recipients, without creating bookings."""
import base64
import binascii
import io
import re
import zipfile
from datetime import datetime, time
from zoneinfo import ZoneInfo

CHILE = ZoneInfo('America/Santiago')


def read_skedu_contacts(encoded, include_unconfirmed=False):
    from openpyxl import load_workbook
    if not isinstance(encoded,str) or len(encoded)>4_000_000:
        raise ValueError('Selecciona un Excel XLSX de hasta 3 MB.')
    try:
        raw=base64.b64decode(encoded,validate=True)
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries=archive.infolist()
            if len(entries)>200 or sum(item.file_size for item in entries)>20_000_000:
                raise ValueError('El Excel supera el tamaño permitido al descomprimirlo.')
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    except (binascii.Error,zipfile.BadZipFile,KeyError,OSError) as error:
        raise ValueError('No se pudo leer el Excel XLSX.') from error
    try:
        if len(workbook.worksheets)<2:raise ValueError('El Excel debe contener la segunda hoja de reservas.')
        sheet=workbook.worksheets[1]
        # Ignore formatting that misleadingly advertises huge worksheet dimensions.
        sheet.reset_dimensions()
        headers=None;records={};skipped=0;unconfirmed=0
        required={'Fecha','Box','Inicio','Término','Cliente','Correo','Confirmación','ID reserva'}
        for number,row in enumerate(sheet.iter_rows(values_only=True),1):
            if number>5005:raise ValueError('El Excel permite hasta 5.000 reservas.')
            if headers is None:
                columns={str(value).strip():index for index,value in enumerate(row) if value is not None}
                if required.issubset(columns):headers=columns
                elif number>=20:raise ValueError('No se encontraron las columnas de reservas en la segunda hoja.')
                continue
            if not any(value is not None for value in row):continue
            def cell(key):
                index=headers[key]
                return row[index] if index<len(row) else None
            day=cell('Fecha');email=str(cell('Correo') or '').strip().lower()
            name=str(cell('Cliente') or '').strip();reference=str(cell('ID reserva') or '').strip()
            status=str(cell('Confirmación') or '').strip().casefold()
            if status not in ('confirmada','no confirmada','cancelada','cancelado','anulada','anulado'):
                raise ValueError(f'Fila {number}: estado de confirmación desconocido.')
            if not isinstance(day,datetime) or day.year<2000 or day.year>2100:
                raise ValueError(f'Fila {number}: fecha de reserva inválida.')
            if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email) or len(email)>200 or not 2<=len(name)<=100 or not 1<=len(reference)<=200:
                raise ValueError(f'Fila {number}: revisa nombre, correo e identificador de reserva.')
            box={'box 1':1,'box 2':2,'box infantil':3}.get(str(cell('Box') or '').strip().casefold())
            start,end=cell('Inicio'),cell('Término')
            if not box or not isinstance(start,time) or not isinstance(end,time) or end<=start:
                raise ValueError(f'Fila {number}: box u horario inválido.')
            active=status=='confirmada' or (status=='no confirmada' and include_unconfirmed)
            record={'reference':reference,'date':day.date().isoformat(),'name':name,'email':email,'box':box,
                    'start':datetime.combine(day.date(),start,CHILE).isoformat(timespec='seconds'),
                    'end':datetime.combine(day.date(),end,CHILE).isoformat(timespec='seconds'),
                    'original_status':status,'active':int(active)}
            if reference in records:
                if records[reference]!=record:raise ValueError(f'Fila {number}: identificador repetido con datos diferentes.')
                continue
            records[reference]=record
            if not active:skipped+=1
            elif status=='no confirmada':unconfirmed+=1
        if headers is None or not records:raise ValueError('La segunda hoja no contiene reservas para importar.')
        return list(records.values()),{'skipped':skipped,'unconfirmed':unconfirmed,'sheet':sheet.title}
    finally:workbook.close()
